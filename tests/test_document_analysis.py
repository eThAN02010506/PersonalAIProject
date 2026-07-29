import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

from qwopus_agent.analysis import AnalysisResult, analyze_uploaded_file
from qwopus_agent.documents import mineru, parse_document, save_uploaded_bytes
from qwopus_agent.documents.mineru import MinerUResult
from qwopus_agent.services.analysis_service import combine_analysis_results


class DocumentAnalysisTests(unittest.TestCase):
    def test_parse_text_document_to_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "notes.txt"
            path.write_text("Hello\nWorld", encoding="utf-8")

            parsed = parse_document(path)

            self.assertEqual(parsed.markdown, "Hello\nWorld")
            self.assertEqual(parsed.metadata["source_type"], "text")

    def test_parse_pdf_prefers_mineru_when_available(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.pdf"
            path.write_bytes(b"%PDF placeholder")

            with patch("qwopus_agent.documents.parser.parse_document_with_mineru") as parse_mineru:
                parse_mineru.return_value = MinerUResult(
                    markdown="# MinerU Markdown",
                    output_path=Path(tmpdir) / "sample.md",
                    command="/usr/local/bin/mineru",
                )

                parsed = parse_document(path)

            self.assertEqual(parsed.markdown, "# MinerU Markdown")
            self.assertEqual(parsed.metadata["parser"], "mineru")
            self.assertEqual(parsed.metadata["mineru_output_path"], str(Path(tmpdir) / "sample.md"))

    def test_parse_docx_prefers_mineru_when_available(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.docx"
            path.write_bytes(b"docx placeholder")

            with patch("qwopus_agent.documents.parser.parse_document_with_mineru") as parse_mineru:
                parse_mineru.return_value = MinerUResult(
                    markdown="# MinerU DOCX Markdown",
                    output_path=Path(tmpdir) / "sample.md",
                    command="/usr/local/bin/mineru",
                )

                parsed = parse_document(path)

            self.assertEqual(parsed.markdown, "# MinerU DOCX Markdown")
            self.assertEqual(parsed.metadata["source_type"], "docx")
            self.assertEqual(parsed.metadata["parser"], "mineru")

    def test_parse_image_uses_mineru_ocr(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "scan.png"
            path.write_bytes(b"image placeholder")

            with patch("qwopus_agent.documents.parser.parse_document_with_mineru") as parse_mineru:
                parse_mineru.return_value = MinerUResult(
                    markdown="# OCR Text",
                    output_path=Path(tmpdir) / "scan.md",
                    command="mineru -b pipeline",
                )

                # 原因：图片不能依赖 PDF/DOCX 的文本回退路径。
                # 作用：验证图片始终通过 MinerU OCR 进入统一 Markdown 分析链路。
                parsed = parse_document(path)

            self.assertEqual(parsed.markdown, "# OCR Text")
            self.assertEqual(parsed.metadata["source_type"], "image")
            self.assertEqual(parsed.metadata["parser"], "mineru")

    def test_mineru_command_prefers_vendor_source(self) -> None:
        with patch.object(mineru, "VENDOR_MINERU_DIR", Path("vendor/MinerU")):
            command = mineru._build_mineru_command()

        self.assertIn("-m", command.args)
        self.assertIn("mineru.cli.client", command.args)
        self.assertIn("vendor/MinerU", command.env["PYTHONPATH"])

    def test_mineru_uses_pipeline_ocr_for_images(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            image_path = Path(tmpdir) / "scan.jpeg"
            image_path.write_bytes(b"image placeholder")
            markdown_path = Path(tmpdir) / "scan.md"
            markdown_path.write_text("OCR result", encoding="utf-8")

            # 原因：图片 real-case 证明 auto 可能只输出图片引用。
            # 作用：锁定 PNG/JPEG 必须使用 pipeline + ocr 的命令契约。
            with (
                patch.object(
                    mineru.subprocess,
                    "run",
                    return_value=SimpleNamespace(returncode=0, stderr=""),
                ) as run,
                patch.object(mineru, "_find_generated_markdown", return_value=markdown_path),
            ):
                mineru.parse_document_with_mineru(
                    image_path,
                    output_root=Path(tmpdir) / "output",
                )

            command = run.call_args.args[0]
            self.assertEqual(command[command.index("-b") + 1], "pipeline")
            self.assertEqual(command[-2:], ["-m", "ocr"])
            self.assertNotEqual(
                Path(command[command.index("-o") + 1]),
                Path(tmpdir) / "output",
            )

    def test_mineru_timeout_uses_the_parser_fallback_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "slow.pdf"
            path.write_bytes(b"%PDF placeholder")

            with (
                patch.object(
                    mineru.subprocess,
                    "run",
                    side_effect=mineru.subprocess.TimeoutExpired("mineru", 300),
                ),
                self.assertRaisesRegex(mineru.MinerUUnavailableError, "could not complete"),
            ):
                # 原因：TimeoutExpired 过去绕过 PDF/DOCX 的回退异常合同。
                # 作用：锁定超时也被标准化，调用方可继续使用 pypdf/python-docx。
                mineru.parse_document_with_mineru(
                    path,
                    output_root=Path(tmpdir) / "output",
                )

    def test_analyze_csv_returns_schema_and_sample_without_full_llm_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sales.csv"
            pd.DataFrame(
                {
                    "region": ["East", "West", "East"],
                    "revenue": [10, 20, 30],
                }
            ).to_csv(path, index=False)

            result = analyze_uploaded_file(path)

            self.assertIn("Spreadsheet Analysis", result.markdown_summary)
            self.assertIn("csv_schema", result.tables)
            self.assertIn("csv_sample", result.tables)
            self.assertEqual(result.metadata["sheets"]["csv"]["rows"], 3)
            self.assertIn("csv_schema", result.markdown_document)

    def test_analyze_excel_returns_local_data_insights(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sales.xlsx"
            pd.DataFrame(
                {
                    "region": ["East", "West", "East", None],
                    "revenue": [10, 20, 30, 40],
                }
            ).to_excel(path, index=False)

            result = analyze_uploaded_file(path, user_question="分析收入")

            self.assertIn("Sheet1_missing_summary", result.tables)
            self.assertIn("Sheet1_numeric_summary", result.tables)
            self.assertIn("Sheet1_categorical_summary", result.tables)
            self.assertIn("numeric_columns", result.metadata["sheets"]["Sheet1"])
            self.assertIn("Sheet1_categorical_summary", result.markdown_document)

    def test_analyze_excel_detects_header_below_title_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "messy_sales.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Report"
            sheet.append(["Sales Report 2026"])
            sheet.append([None, None])
            sheet.append(["region", "revenue"])
            sheet.append(["East", 10])
            sheet.append(["West", 20])
            workbook.save(path)

            result = analyze_uploaded_file(path)

            sheet_metadata = result.metadata["sheets"]["Report"]
            self.assertEqual(sheet_metadata["column_names"], ["region", "revenue"])
            self.assertEqual(sheet_metadata["header_row"], 3)
            self.assertIn("revenue", sheet_metadata["numeric_columns"])

    def test_analyze_excel_handles_sheets_without_numeric_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "text_only.xlsx"
            pd.DataFrame(
                {
                    "name": ["Alice", "Bob"],
                    "note": ["ready", "blocked"],
                }
            ).to_excel(path, index=False)

            result = analyze_uploaded_file(path)

            self.assertIn("Sheet1_schema", result.tables)
            self.assertIn("Sheet1_categorical_summary", result.tables)
            self.assertNotIn("Sheet1_numeric_summary", result.tables)

    def test_analyze_form_like_excel_extracts_key_value_pairs(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "character_sheet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "人物卡"
            sheet.append(["调查员信息", None, None])
            sheet.append(["姓名", "L'hopital", None])
            sheet.append(["职业", "教授", None])
            sheet.append(["年龄", 50, None])
            workbook.save(path)

            result = analyze_uploaded_file(path)
            form_summary = result.tables["人物卡_form_summary"]

            pairs = {row["key"]: row["value"] for row in form_summary.to_dict(orient="records")}
            self.assertEqual(pairs["姓名"], "L'hopital")
            self.assertEqual(pairs["职业"], "教授")
            self.assertIn("人物卡_form_summary", result.markdown_document)

    def test_combine_analysis_results_keeps_multiple_file_contexts(self) -> None:
        result_a = AnalysisResult(
            markdown_summary="# A",
            tables={"metadata": pd.DataFrame([{"key": "a", "value": 1}])},
            metadata={"source_type": "text"},
            markdown_document="alpha",
        )
        result_b = AnalysisResult(
            markdown_summary="# B",
            tables={"metadata": pd.DataFrame([{"key": "b", "value": 2}])},
            metadata={"source_type": "text"},
            markdown_document="beta",
        )

        combined = combine_analysis_results(
            [
                ("a.txt", result_a),
                ("b.txt", result_b),
            ]
        )

        self.assertEqual(combined.metadata["file_count"], 2)
        self.assertIn("# File: a.txt", combined.markdown_document)
        self.assertIn("# File: b.txt", combined.markdown_document)
        self.assertIn("a::metadata", combined.tables)
        self.assertIn("b::metadata", combined.tables)

    def test_analysis_result_can_carry_llm_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "notes.txt"
            path.write_text("Project Alpha increased revenue.", encoding="utf-8")

            result = analyze_uploaded_file(path, user_question="总结")

            enriched = type(result)(
                markdown_summary=result.markdown_summary,
                tables=result.tables,
                metadata=result.metadata,
                markdown_document=result.markdown_document,
                llm_analysis="这是一份关于 Project Alpha 收入增长的文档。",
            )

            self.assertEqual(enriched.llm_analysis, "这是一份关于 Project Alpha 收入增长的文档。")

    def test_save_uploaded_bytes_uses_upload_directory(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            stored = save_uploaded_bytes(
                filename="../unsafe.txt",
                content=b"data",
                upload_dir=Path(tmpdir),
            )

            self.assertEqual(stored.original_name, "unsafe.txt")
            self.assertTrue(stored.path.exists())
            self.assertEqual(stored.path.read_bytes(), b"data")
