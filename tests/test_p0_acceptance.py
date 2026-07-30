import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from qwopus_agent.analysis import analyze_uploaded_file
from qwopus_agent.evaluation import RetrievalBenchmarkCase, evaluate_retrieval
from tests.minirag_fakes import make_test_minirag


class P0AcceptanceTests(unittest.TestCase):
    def test_large_multilingual_document_keeps_tail_evidence_after_restart(self) -> None:
        with TemporaryDirectory() as tmpdir:
            storage_path = Path(tmpdir) / "documents.jsonl"
            memory = make_test_minirag(storage_path)
            sections = "\n\n".join(
                f"# Chapter {number}\n"
                f"This is bounded section {number}. 第 {number} 章保留独立结构和来源。"
                for number in range(1, 81)
            )
            tail_fact = "最终验收标识 ORBIT-DELTA-731 属于第八十章。"
            memory.insert(f"# File: large-guide.md\n\n{sections}\n\n{tail_fact}")

            reloaded = make_test_minirag(storage_path)
            results = reloaded.search("ORBIT DELTA 731 第八十章", min_relevance=0.25)
            report = evaluate_retrieval(
                RetrievalBenchmarkCase(
                    name="large-document-tail",
                    query="ORBIT DELTA 731 第八十章",
                    expected_sources=("large-guide.md",),
                ),
                results,
            )

        self.assertTrue(report.passed, report)
        self.assertIn(tail_fact, "\n".join(results))

    def test_multitable_non_english_workbook_preserves_each_region(self) -> None:
        with TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "区域报表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "运营数据"
            sheet.append(["地区", "收入"])
            sheet.append(["华东", 120])
            sheet.append(["华西", 90])
            sheet.append([])
            sheet.append(["团队", "工单"])
            sheet.append(["Alpha", 4])
            sheet.append(["Beta", 7])
            workbook.save(path)

            result = analyze_uploaded_file(path, user_question="比较收入与工单")

        profile = result.metadata["workbook_profile"]
        self.assertEqual(profile["sheets"][0]["kind"], "multi_table")
        self.assertIn("运营数据::table_2_schema", result.tables)
        self.assertIn("运营数据::table_2_sample", result.tables)


if __name__ == "__main__":
    unittest.main()
